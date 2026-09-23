#!/usr/bin/env python3
"""Chuyển đổi báo cáo giải trình mua sắm (Markdown) sang bản Word (.docx)
trình bày theo phong cách báo cáo chuyên nghiệp.

Định dạng:
  - Khổ A4 dọc (riêng bảng >= 8 cột tự động sang trang khổ ngang);
  - Font Times New Roman 13 pt, giãn dòng 1.15;
  - Khối thông tin đầu báo cáo và bảng tiêu chí từng thiết bị: cột nhãn đổ màu;
  - Bảng dữ liệu: hàng tiêu đề đổ màu xanh đậm, chữ trắng, zebra row;
  - Tiêu đề các cấp có màu và đường kẻ chân; tiêu đề trang + số trang tự động.

Cách dùng:  python3 tools/md2docx.py [file.md] [file.docx]
"""
import os
import re
import sys

_LIBS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".pylibs")
if os.path.isdir(_LIBS) and _LIBS not in sys.path:
    sys.path.insert(0, _LIBS)

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = sys.argv[1] if len(sys.argv) > 1 else "bao-cao-giai-trinh-mua-sam-phan-bo-thiet-bi.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else "Bao-cao-giai-trinh-mua-sam-phan-bo-thiet-bi.docx"

DARK = "1F4E79"      # xanh đậm tiêu đề / đầu bảng
MID = "2E74B5"       # xanh vừa cho heading cấp 2
LABEL = "DCE6F1"     # nền cột nhãn
ZEBRA = "F2F7FC"     # nền dòng chẵn
INFO = "EDF3FA"      # nền khối thông tin
GRAY = "595959"

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
SEP_CELL = re.compile(r"^:?-{2,}:?$")

WIDTHS = {
    ("Stt", "Danh mục hàng hóa"): [0.9, 2.6, 2.2, 1.7, 1.0, 1.6, 1.9, 6.2, 1.0, 0.9, 6.1],
    ("STT", "Tài liệu"): [1.2, 6.4, 9.8],
    ("#", "Thiết bị"): [1.0, 7.2, 9.2],
    ("Môn học (mã – tên)", "Tiết TH"): [8.8, 2.6, 6.0],
    ("Môn học (mã – tên)", "TH (TC)"): [8.8, 2.6, 6.0],
}


# ---------------------------------------------------------------- tiện ích OOXML
def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def para_border(par, edges=("bottom",), color=DARK, sz=8, space=3):
    p_pr = par._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        p_bdr.append(el)
    p_pr.append(p_bdr)


def fixed_layout(table):
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(lay)


def set_widths(table, widths):
    fixed_layout(table)
    for row in table.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = Cm(w)


def cell_margins(table, left=80, right=80, top=40, bottom=40):
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("left", left), ("right", right), ("top", top), ("bottom", bottom)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def repeat_header(table):
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_page_field(par):
    par.add_run("Trang ")
    r = par.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1)
    r._r.append(instr)
    r._r.append(f2)


def add_inline(par, text, base_bold=False, size=None, color=None):
    for tok in INLINE.split(text.replace("<br>", "\n")):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = par.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run = par.add_run(tok[1:-1])
            run.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = par.add_run(tok[1:-1])
            run.font.name = "Consolas"
        else:
            run = par.add_run(tok)
            if base_bold:
                run.bold = True
        if size is not None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = RGBColor.from_string(color)


# ---------------------------------------------------------------- bảng biểu
def build_table(doc, rows):
    hdr = rows[0]
    label_table = len(hdr) == 2 and hdr[0].strip() in ("", "Tiêu chí", "Hạng mục")
    if label_table:
        rows = rows[1:]
    n_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.autofit = False
    cell_margins(table)

    if label_table:
        set_widths(table, [4.6, 12.8] if n_cols == 2 else [None] * n_cols)
        for i, row in enumerate(rows):
            for j, text in enumerate(row):
                cell = table.cell(i, j)
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(3)
                par.paragraph_format.space_before = Pt(3)
                add_inline(par, text, size=12)
                if j == 0:
                    shade(cell, LABEL)
                    for r in par.runs:
                        r.bold = True
        return table

    repeat_header(table)
    key = (hdr[0].strip(), hdr[1].strip()) if n_cols >= 2 else None
    if n_cols >= 8:
        widths = WIDTHS.get(key, [26.1 / n_cols] * n_cols)
        font_size = 8.5
    else:
        widths = WIDTHS.get(key, [17.4 / n_cols] * n_cols)
        font_size = 10.5
    set_widths(table, widths)

    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.space_before = Pt(2)
            if i == 0:
                shade(cell, DARK)
                add_inline(par, text, base_bold=True, size=font_size + 0.5, color="FFFFFF")
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if i % 2 == 0:
                    shade(cell, ZEBRA)
                add_inline(par, text, size=font_size)
    return table


def new_section(doc, landscape):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    return sec


def decorate_section(sec, running_head):
    sec.different_first_page_header_footer = True
    for footer in (sec.footer, sec.first_page_footer):
        par = footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_field(par)
        for r in par.runs:
            r.font.size = Pt(10)
            r.font.name = "Times New Roman"
            r.font.color.rgb = RGBColor.from_string(GRAY)
    hpar = sec.header.paragraphs[0]
    hpar.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hpar.add_run(running_head)
    run.italic = True
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"
    run.font.color.rgb = RGBColor.from_string(GRAY)


# ---------------------------------------------------------------- thân chương trình
def main():
    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in (
        ("Title", 19, DARK),
        ("Heading 1", 15, DARK),
        ("Heading 2", 13.5, MID),
        ("Heading 3", 13, MID),
        ("List Bullet", 13, "000000"),
        ("List Bullet 2", 12.5, "000000"),
        ("List Number", 13, "000000"),
    ):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    first = doc.sections[0]
    first.orientation = WD_ORIENT.PORTRAIT
    first.page_width, first.page_height = Cm(21.0), Cm(29.7)
    first.left_margin = first.right_margin = Cm(1.8)
    first.top_margin = first.bottom_margin = Cm(1.8)

    running_head = "Báo cáo giải trình mua sắm & phân bổ thiết bị – Khoa Vật lý – Vật lý Kỹ thuật"
    decorate_section(first, running_head)
    in_landscape = False

    doc.core_properties.title = "Báo cáo giải trình mua sắm & phân bổ thiết bị"
    doc.core_properties.author = "Khoa Vật lý – Vật lý Kỹ thuật"
    doc.core_properties.comments = "Bản v1.3 – sinh tự động từ Markdown bằng tools/md2docx.py"

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(SEP_CELL.match(c) for c in cells):
                    block.append(cells)
                i += 1
            n_cols = max(len(r) for r in block)
            block = [r + [""] * (n_cols - len(r)) for r in block]
            need_landscape = n_cols >= 8
            if need_landscape != in_landscape:
                sec = new_section(doc, need_landscape)
                decorate_section(sec, running_head)
                in_landscape = need_landscape
            build_table(doc, block)
            doc.add_paragraph()
            continue

        if stripped.startswith("# "):
            par = doc.add_heading(level=0)
            add_inline(par, stripped[2:])
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para_border(par, ("bottom",), DARK, 12, 6)
            par.paragraph_format.space_after = Pt(14)
        elif stripped.startswith("## "):
            text = stripped[3:]
            if text[:4] in ("III.", "V. P", "VI. "):
                doc.add_page_break()
            par = doc.add_heading(level=1)
            add_inline(par, text)
            para_border(par, ("bottom",), MID, 6, 3)
            par.paragraph_format.space_before = Pt(14)
            par.paragraph_format.space_after = Pt(8)
        elif stripped.startswith("### "):
            par = doc.add_heading(level=2)
            add_inline(par, stripped[4:])
            par.paragraph_format.space_before = Pt(12)
            par.paragraph_format.space_after = Pt(6)
        elif stripped == "---":
            pass
        elif stripped.startswith("> "):
            par = doc.add_paragraph()
            add_inline(par, stripped[2:], size=12, color=GRAY)
            for run in par.runs:
                run.italic = True
            par.paragraph_format.left_indent = Cm(0.6)
            para_border(par, ("left",), MID, 18, 8)
        elif stripped.startswith("- "):
            par = doc.add_paragraph(style="List Bullet")
            add_inline(par, stripped[2:])
        elif re.match(r"^ {2,}- ", lines[i]):
            par = doc.add_paragraph(style="List Bullet 2")
            add_inline(par, stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            par = doc.add_paragraph(style="List Number")
            add_inline(par, re.sub(r"^\d+\. ", "", stripped))
        elif stripped:
            par = doc.add_paragraph()
            add_inline(par, stripped)
        i += 1

    doc.save(OUT)
    print(f"Đã tạo {OUT}: {len(doc.paragraphs)} đoạn, {len(doc.tables)} bảng, {len(doc.sections)} section")


if __name__ == "__main__":
    main()
